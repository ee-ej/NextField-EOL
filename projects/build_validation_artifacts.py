from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "source" / "Flex Asset EOL and Site Risk Evaluation Tool Ver 3.0 3.6.25.xlsx"
REPORT_DIR = ROOT / "reports" / "validation"


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def normalize_text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def find_header(ws: Any, label: str, max_scan_row: int = 80) -> tuple[int, int]:
    for row_number in range(1, min(ws.max_row, max_scan_row) + 1):
        for column_number in range(1, ws.max_column + 1):
            if normalize_text(ws.cell(row_number, column_number).value) == label:
                return row_number, column_number
    raise ValueError(f"Could not find {label!r} on {ws.title!r}")


def valid_asset_id(value: Any) -> bool:
    return is_number(value) and value > 0


def valid_site_no(value: Any, known_sites: set[str] | None = None) -> bool:
    text = normalize_text(value)
    if not text or text in {"-", "N/A", "0"}:
        return False
    if known_sites is not None:
        return text in known_sites
    return bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9.\-]*", text))


def collect_key_rows(ws: Any, header_row: int, key_column: int, predicate: Any) -> tuple[list[dict[str, Any]], list[int]]:
    valid_rows: list[dict[str, Any]] = []
    excluded_rows: list[int] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        key = values[key_column - 1] if key_column <= len(values) else None
        if predicate(key):
            valid_rows.append({"row": row_number, "key": json_value(key)})
        else:
            excluded_rows.append(row_number)
    return valid_rows, excluded_rows


def write_csv(path: Path, rows: Iterable[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_kpi_manifest(wb: Any) -> dict[str, Any]:
    site_ws = wb["Entry- Site Data"]
    site_header, site_key_column = find_header(site_ws, "Site No.")
    site_rows, site_excluded = collect_key_rows(
        site_ws, site_header, site_key_column, lambda value: valid_site_no(value)
    )
    known_sites = {normalize_text(row["key"]) for row in site_rows}

    asset_ws = wb["Entry- Eq Database"]
    asset_header, asset_key_column = find_header(asset_ws, "AssetId")
    asset_rows, asset_excluded = collect_key_rows(
        asset_ws, asset_header, asset_key_column, valid_asset_id
    )

    scored_ws = wb["Output- Asset Scores"]
    scored_header, scored_key_column = find_header(scored_ws, "AssetId")
    scored_rows, scored_excluded = collect_key_rows(
        scored_ws, scored_header, scored_key_column, valid_asset_id
    )

    scored_site_column = find_header(scored_ws, "SiteNo")[1]
    scored_context_rows = []
    for row_number, values in enumerate(
        scored_ws.iter_rows(min_row=scored_header + 1, values_only=True), start=scored_header + 1
    ):
        asset_id = values[scored_key_column - 1] if scored_key_column <= len(values) else None
        site_no = values[scored_site_column - 1] if scored_site_column <= len(values) else None
        if valid_asset_id(asset_id) and valid_site_no(site_no, known_sites):
            scored_context_rows.append(row_number)

    site_score_ws = wb["Output- Site Scores"]
    site_score_header, site_score_key_column = find_header(site_score_ws, "Site No.")
    site_score_rows, site_score_excluded = collect_key_rows(
        site_score_ws,
        site_score_header,
        site_score_key_column,
        lambda value: valid_site_no(value, known_sites),
    )

    risk_ws = wb["Output- Site Risk Scores"]
    risk_header, risk_key_column = find_header(risk_ws, "Site No.")
    risk_rows, risk_excluded = collect_key_rows(
        risk_ws,
        risk_header,
        risk_key_column,
        lambda value: valid_site_no(value, known_sites),
    )

    def profile(rows: list[dict[str, Any]], excluded: list[int], header: int, max_row: int) -> dict[str, Any]:
        keys = [str(row["key"]) for row in rows]
        duplicates = {key: count for key, count in Counter(keys).items() if count > 1}
        return {
            "header_row": header,
            "first_valid_row": rows[0]["row"] if rows else None,
            "last_valid_row": rows[-1]["row"] if rows else None,
            "valid_row_count": len(rows),
            "distinct_key_count": len(set(keys)),
            "duplicate_keys": duplicates,
            "excluded_row_count": len(excluded),
            "used_range_max_row": max_row,
        }

    source_asset_keys = {str(row["key"]) for row in asset_rows}
    scored_asset_keys = {str(row["key"]) for row in scored_rows}
    return {
        "workbook": str(WORKBOOK_PATH),
        "validity_rules": {
            "asset": "numeric AssetId greater than zero",
            "site": "nonblank SiteNo present in bounded Entry- Site Data population",
            "output": "valid nonzero business key; output rows are derived facts",
        },
        "kpis": {
            "sites_in_portfolio": len(known_sites),
            "assets_in_source": len(source_asset_keys),
            "equipment_rows": len(asset_rows),
            "scored_assets": len(scored_asset_keys),
            "scored_rows_with_site_context": len(scored_context_rows),
            "source_assets_absent_from_scored_output": sorted(source_asset_keys - scored_asset_keys),
        },
        "tables": {
            "Entry- Site Data": profile(site_rows, site_excluded, site_header, site_ws.max_row),
            "Entry- Eq Database": profile(asset_rows, asset_excluded, asset_header, asset_ws.max_row),
            "Output- Asset Scores": profile(scored_rows, scored_excluded, scored_header, scored_ws.max_row),
            "Output- Site Scores": profile(site_score_rows, site_score_excluded, site_score_header, site_score_ws.max_row),
            "Output- Site Risk Scores": profile(risk_rows, risk_excluded, risk_header, risk_ws.max_row),
        },
    }


def extract_spof(wb: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    ws = wb["Entry- Site SPOF Data"]
    group_row = 24
    date_row = 26
    lead_row = 27
    answer_row = 28
    question_start = 30
    group_starts: list[int] = []

    for column_number in range(7, ws.max_column + 1):
        if normalize_text(ws.cell(group_row, column_number).value) and normalize_text(
            ws.cell(answer_row, column_number).value
        ) == "Yes":
            group_starts.append(column_number)

    rows: list[dict[str, Any]] = []
    question_count = 0
    for question_row in range(question_start, ws.max_row + 1):
        question_id = ws.cell(question_row, 1).value
        question_text = normalize_text(ws.cell(question_row, 2).value)
        if not is_number(question_id) or not question_text:
            continue
        question_count += 1
        for start_column in group_starts:
            site_no = normalize_text(ws.cell(group_row, start_column).value)
            markers = [
                normalize_text(ws.cell(question_row, start_column + offset).value)
                for offset in range(3)
            ]
            selected = [
                label
                for label, marker in zip(("Yes", "No", "N/A"), markers)
                if marker.lower() == "x"
            ]
            status = "Valid"
            if len(selected) > 1:
                status = "MultipleAnswers"
            elif not site_no:
                status = "UnmappedSite"
            answer = selected[0] if len(selected) == 1 else "Blank"
            if len(selected) > 1:
                answer = "|".join(selected)
            rows.append(
                {
                    "SiteNo": site_no,
                    "SPOFQuestionId": str(question_id),
                    "SPOFQuestion": question_text,
                    "AnswerValue": answer,
                    "OriginalMarker": "|".join(marker for marker in markers if marker),
                    "DetailedComments": normalize_text(ws.cell(question_row, start_column + 3).value),
                    "SourceSheet": ws.title,
                    "SourceColumnStart": start_column,
                    "SourceColumnEnd": start_column + 3,
                    "SourceRows": f"{group_row},{date_row},{lead_row},{answer_row},{question_row}",
                    "SourceDate": json_value(ws.cell(date_row, start_column).value),
                    "Lead": normalize_text(ws.cell(lead_row, start_column).value),
                    "NormalizationStatus": status,
                }
            )

    summary = {
        "sheet": ws.title,
        "group_row": group_row,
        "question_start_row": question_start,
        "group_count": len(group_starts),
        "question_count": question_count,
        "observation_count": len(rows),
        "group_sites": sorted({row["SiteNo"] for row in rows if row["SiteNo"]}),
        "status_counts": dict(Counter(row["NormalizationStatus"] for row in rows)),
    }
    return rows, summary


def main() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    manifest = build_kpi_manifest(wb)
    spof_rows, spof_summary = extract_spof(wb)
    manifest["spof"] = spof_summary

    (REPORT_DIR / "bounded_kpi_manifest.json").write_text(
        json.dumps(manifest, indent=2, default=json_value), encoding="utf-8"
    )
    write_csv(
        REPORT_DIR / "spof_observations.csv",
        spof_rows,
        [
            "SiteNo",
            "SPOFQuestionId",
            "SPOFQuestion",
            "AnswerValue",
            "OriginalMarker",
            "DetailedComments",
            "SourceSheet",
            "SourceColumnStart",
            "SourceColumnEnd",
            "SourceRows",
            "SourceDate",
            "Lead",
            "NormalizationStatus",
        ],
    )
    print(json.dumps(manifest, indent=2, default=json_value))
    print(f"Wrote artifacts to {REPORT_DIR}")


if __name__ == "__main__":
    main()
